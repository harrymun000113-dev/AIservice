"""엑셀(AI 서비스 조사 통합표) -> js/services.js 변환 스크립트.
사용법: python tools/xlsx_to_services.py "엑셀파일경로.xlsx"
"""
import json, re, sys, os
from openpyxl import load_workbook

src = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\user\Downloads\AI_서비스_조사_통합.xlsx(최종).xlsx"
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "js", "services.js")

CATS = ["리서치", "글쓰기", "이미지 생성", "영상", "음성·음악", "녹음·회의록", "시각화·PPT", "웹·UI/UX 디자인", "AI 만화·스토리보드", "업무 자동화"]
SLUG = {"리서치": "research", "글쓰기": "writing", "이미지 생성": "image", "영상": "video", "음성·음악": "audio",
        "녹음·회의록": "meeting", "시각화·PPT": "slides", "웹·UI/UX 디자인": "webdesign",
        "AI 만화·스토리보드": "comic", "업무 자동화": "automation"}
GROUP_OVERRIDE = {"클로바노트": "clovanote", "라이너 (Liner)": "liner"}
# 가장 저렴한 유료 플랜 월 가격(USD) 수동 보정 (원화·유로·연 표기 등)
MIN_OVERRIDE = {10: 7.99, 29: 12.0, 21: round(20000 / 1400, 2), 63: round(9900 / 1400, 2), 66: round(24 * 1.08, 2)}
NO_PRICE = {8, 11, 16, 17, 33, 38, 39, 42, 44, 48, 53}  # 무료/미공개/금액 없음

def norm_price(p):
    p = (p or "").strip()
    if p.startswith("무료+유료"): return "무료+유료"
    if p.startswith("무료"): return "무료"
    if p.startswith("유료"): return "유료"
    return "무료+유료"

def bullets(t):
    return [re.sub(r"^[•·\-\s]+", "", x).strip() for x in (t or "").split("\n") if x.strip()]

def group_of(name):
    if name in GROUP_OVERRIDE: return GROUP_OVERRIDE[name]
    m = re.search(r"\(([A-Za-z0-9 .]+)\)", name)
    base = m.group(1) if m else name
    return re.sub(r"[^a-z0-9가-힣]", "", base.lower())

def min_paid(no, note):
    if no in MIN_OVERRIDE: return MIN_OVERRIDE[no]
    if no in NO_PRICE: return None
    lines = [l for l in (note or "").split("\n") if l.strip() and not l.strip().startswith("※")]
    if not lines: return None
    m = re.search(r"\$\s?([\d.,]+)", lines[0])
    return float(m.group(1).replace(",", "")) if m else None

ws = load_workbook(src, data_only=True).active
items = []
for r in ws.iter_rows(min_row=5, values_only=True):
    if not isinstance(r[0], (int, float)): continue
    no = int(r[0]); cat = r[1]; name = r[2]
    if cat not in CATS: raise SystemExit("알 수 없는 분야: %r" % cat)
    note = r[9] or ""
    note_lines = [l.strip() for l in note.split("\n") if l.strip()]
    plans = [l for l in note_lines if not l.startswith("※")]
    notes = [l.lstrip("※ ").strip() for l in note_lines if l.startswith("※")]
    tried = any("미체험" in n for n in notes)
    notes = [n for n in notes if "미체험" not in n]
    items.append({
        "id": "%s-%s" % (group_of(name), SLUG[cat]),
        "group": group_of(name),
        "no": no,
        "name": name,
        "category": cat,
        "desc": r[3],
        "url": r[4],
        "pros": bullets(r[5]),
        "cons": bullets(r[6]),
        "use": r[7],
        "price": norm_price(r[8]),
        "priceDetail": (r[8] or "").strip(),
        "plans": plans,
        "priceNotes": notes,
        "minPaidUSD": min_paid(no, note),
        "researchedOnly": tried,
    })
ids = [i["id"] for i in items]
assert len(ids) == len(set(ids)), "id 중복"
with open(out, "w", encoding="utf-8") as f:
    f.write("// 자동 생성 파일 (tools/xlsx_to_services.py). 직접 수정하지 말고 엑셀을 고친 뒤 다시 생성하세요.\n")
    f.write("const CATEGORIES = %s;\n" % json.dumps(CATS, ensure_ascii=False))
    f.write("const SERVICES = %s;\n" % json.dumps(items, ensure_ascii=False, indent=1))
# 순수 데이터 파일(JSON·CSV)도 함께 저장: 다른 곳에서 재사용하거나 엑셀로 열 때 사용
import csv
ddir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data")
os.makedirs(ddir, exist_ok=True)
with open(os.path.join(ddir, "ai-services.json"), "w", encoding="utf-8") as f:
    json.dump({"updated": "2026-10-07", "categories": CATS, "services": items}, f, ensure_ascii=False, indent=1)
with open(os.path.join(ddir, "ai-services.csv"), "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["No", "분야", "서비스 이름", "한줄 소개", "주소", "요금", "요금 상세", "장점", "단점", "추천 용도", "가장 저렴한 유료 플랜(월 USD)", "요금 플랜", "비고"])
    for i in items:
        w.writerow([i["no"], i["category"], i["name"], i["desc"], i["url"], i["price"], i["priceDetail"],
                    " / ".join(i["pros"]), " / ".join(i["cons"]), i["use"],
                    "" if i["minPaidUSD"] is None else i["minPaidUSD"], " / ".join(i["plans"]), " / ".join(i["priceNotes"])])
print(len(items), "건 저장 ->", os.path.normpath(out))
for i in items:
    print(i["no"], i["name"], i["price"], i["minPaidUSD"], "(조사만)" if i["researchedOnly"] else "")
